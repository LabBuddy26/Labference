from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

try:
    from openpyxl import Workbook
except ImportError:  # pragma: no cover - baseline environments may omit optional deps
    Workbook = None  # type: ignore[assignment,misc]

from eai4s_eval.dataset_statistics import (
    DEFAULT_WORKBOOK_PATH,
    DatasetStatisticsValidationError,
    _reconcile_simulation_summary,
    build_statistics_payload,
    duration_to_seconds,
    generate_statistics_json,
)


def _dictionary_rows() -> list[list[object]]:
    rows: list[list[object]] = []
    tasks = (
        ("liquid_pour", "液体倾倒"),
        ("test_tube_to_rack", "试管入架"),
        ("centrifuge_tube_to_rack", "离心管入架"),
        ("sample_vial_storage", "样品瓶收纳"),
        ("glass_beaker_pick_place", "玻璃烧杯抓放"),
        ("clear_bottle_transfer", "透明瓶转移"),
    )
    rows.extend([["task", code, name, "", "", True, ""] for code, name in tasks])
    rows.extend(
        [
            ["device", "XARM", "XARM", "real_robot_teleoperation", "single_arm", True, ""],
            ["device", "UR3", "UR3", "real_robot_teleoperation", "single_arm", True, ""],
            ["device", "OPENARM", "OPENARM", "real_robot_teleoperation", "dual_arm", True, ""],
            ["device", "UMI", "UMI", "umi", "dual_arm", True, ""],
            ["simulator", "ISAAC_SIM", "Isaac Sim", "", "", True, ""],
            ["robot", "FRANKA", "Franka", "", "", True, ""],
        ]
    )
    return rows


def _make_workbook(path: Path, physical: list[list[object]], simulation: list[list[object]]) -> None:
    assert Workbook is not None
    workbook = Workbook()
    workbook.active.title = "使用说明"
    physical_sheet = workbook.create_sheet("真机UMI明细")
    physical_sheet.append(
        [
            "batch_id", "device_model", "collection_method", "arm_mode", "task_id",
            "task_name", "raw_task_name", "task_description", "collection_date",
            "sample_count", "valid_duration", "operator_ids", "notes",
        ]
    )
    for row in physical:
        physical_sheet.append(row)
        physical_sheet.cell(physical_sheet.max_row, 11).number_format = "[h]:mm:ss"
    simulation_sheet = workbook.create_sheet("仿真明细")
    simulation_sheet.append(
        [
            "batch_id", "dataset_id", "simulator", "robot_type", "task_id", "task_name",
            "scene_id", "generation_date", "episode_count", "frame_count",
            "simulation_duration", "notes",
        ]
    )
    for row in simulation:
        simulation_sheet.append(row)
        simulation_sheet.cell(simulation_sheet.max_row, 11).number_format = "[h]:mm:ss"
    dictionary_sheet = workbook.create_sheet("配置字典")
    dictionary_sheet.append(
        ["dictionary_type", "code", "name", "collection_method", "arm_mode", "enabled", "notes"]
    )
    for row in _dictionary_rows():
        dictionary_sheet.append(row)
    workbook.create_sheet("统计汇总")
    workbook.save(path)
    workbook.close()


@unittest.skipIf(Workbook is None, "openpyxl is not installed")
class DatasetStatisticsTest(unittest.TestCase):
    def test_excel_duration_conversion(self) -> None:
        self.assertEqual(duration_to_seconds("00:14:27"), 867)
        self.assertEqual(duration_to_seconds(timedelta(minutes=14, seconds=27)), 867)
        self.assertEqual(duration_to_seconds(867 / 86400), 867)

    def test_physical_aggregation_and_operator_union(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "registry.xlsx"
            rows = [
                [
                    "batch-1", "UR3", "real_robot_teleoperation", "single_arm",
                    "centrifuge_tube_to_rack", "离心管入架", "Pick Centrifuge Tube",
                    "description", date(2026, 6, 23), 100, timedelta(seconds=1000),
                    "operator01;operator02", "",
                ],
                [
                    "batch-2", "UR3", "real_robot_teleoperation", "single_arm",
                    "centrifuge_tube_to_rack", "离心管入架", "Pick Centrifuge Tube",
                    "description", date(2026, 7, 6), 160, timedelta(seconds=3555),
                    "operator02", "",
                ],
            ]
            _make_workbook(workbook_path, rows, [])
            payload = build_statistics_payload(workbook_path, validate_summary=False)

        section = payload["real_robot_teleoperation"]
        self.assertEqual(section["totals"], {"batch_count": 2, "sample_count": 260, "duration_seconds": 4555})
        self.assertEqual(len(section["rows"]), 1)
        group = section["rows"][0]
        self.assertEqual(group["collection_date_start"], "2026-06-23")
        self.assertEqual(group["collection_date_end"], "2026-07-06")
        self.assertEqual(group["operator_ids"], ["operator01", "operator02"])

    def test_empty_and_populated_simulation_are_independent(self) -> None:
        physical = [
            [
                "batch-real", "UR3", "real_robot_teleoperation", "single_arm",
                "liquid_pour", "液体倾倒", "pour_water", "description", date(2026, 4, 29),
                41, timedelta(seconds=1698), "operator01", "",
            ]
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            empty_path = Path(temp_dir) / "empty.xlsx"
            _make_workbook(empty_path, physical, [])
            empty_payload = build_statistics_payload(empty_path, validate_summary=False)
            simulation = [
                [
                    "sim-batch-1", "dataset-1", "ISAAC_SIM", "FRANKA", "liquid_pour",
                    "液体倾倒", "scene-a", date(2026, 7, 1), 10, 1000,
                    timedelta(seconds=100), "",
                ],
                [
                    "sim-batch-2", "dataset-2", "ISAAC_SIM", "FRANKA", "liquid_pour",
                    "液体倾倒", "scene-b", date(2026, 7, 3), 20, 3000,
                    timedelta(seconds=250), "",
                ],
            ]
            populated_path = Path(temp_dir) / "populated.xlsx"
            _make_workbook(populated_path, physical, simulation)
            populated_payload = build_statistics_payload(populated_path, validate_summary=False)

        self.assertEqual(empty_payload["simulation"]["status"], "empty")
        self.assertEqual(
            empty_payload["simulation"]["totals"],
            {"dataset_batch_count": 0, "episode_count": 0, "frame_count": 0, "duration_seconds": 0, "scene_count": 0},
        )
        self.assertEqual(
            populated_payload["real_robot_teleoperation"]["totals"],
            empty_payload["real_robot_teleoperation"]["totals"],
        )
        sim_group = populated_payload["simulation"]["rows"][0]
        self.assertEqual(sim_group["dataset_batch_count"], 2)
        self.assertEqual(sim_group["episode_count"], 30)
        self.assertEqual(sim_group["frame_count"], 4000)
        self.assertEqual(sim_group["duration_seconds"], 350)
        self.assertEqual(sim_group["scene_count"], 2)

    def test_summary_must_include_every_nonzero_simulation_group(self) -> None:
        assert Workbook is not None
        formula_book = Workbook()
        formula_sheet = formula_book.active
        cached_book = Workbook()
        cached_sheet = cached_book.active
        headers = [
            "simulator", "robot_type", "task_id", "task_name", "dataset_batch_count",
            "episode_count", "frame_count", "duration", "scene_count", "first_date", "last_date",
        ]
        for sheet in (formula_sheet, cached_sheet):
            sheet.append(["仿真汇总"] + [None] * 10)
            sheet.append(headers)
            sheet.append(["暂无仿真数据"] + [None] * 10)
            sheet.append(["仿真总计", None, None, None, 0, 0, 0, 0, 0, None, None])
        for column in range(5, 12):
            formula_sheet.cell(4, column).value = "=0"
        section = {
            "totals": {
                "dataset_batch_count": 1,
                "episode_count": 10,
                "frame_count": 100,
                "duration_seconds": 20,
                "scene_count": 1,
            },
            "rows": [{
                "simulator": "ISAAC_SIM",
                "robot_type": "FRANKA",
                "task_id": "liquid_pour",
                "task_name": "液体倾倒",
                "dataset_batch_count": 1,
                "episode_count": 10,
                "frame_count": 100,
                "duration_seconds": 20,
                "scene_count": 1,
                "generation_date_start": "2026-07-01",
                "generation_date_end": "2026-07-01",
            }],
        }
        with self.assertRaisesRegex(DatasetStatisticsValidationError, "missing simulation group row"):
            _reconcile_simulation_summary(formula_sheet, cached_sheet, section, 1, 4)

    def test_invalid_operator_and_duplicate_ids_fail_validation(self) -> None:
        base = [
            "batch-1", "UR3", "real_robot_teleoperation", "single_arm", "liquid_pour",
            "液体倾倒", "pour_water", "description", date(2026, 4, 29), 10,
            timedelta(seconds=20), "真实姓名", "",
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            invalid_path = Path(temp_dir) / "invalid.xlsx"
            _make_workbook(invalid_path, [base], [])
            with self.assertRaisesRegex(DatasetStatisticsValidationError, "operatorNN"):
                build_statistics_payload(invalid_path, validate_summary=False)

            duplicate_path = Path(temp_dir) / "duplicate.xlsx"
            valid = [*base]
            valid[11] = "operator01"
            _make_workbook(duplicate_path, [valid, valid], [])
            with self.assertRaisesRegex(DatasetStatisticsValidationError, "duplicated"):
                build_statistics_payload(duplicate_path, validate_summary=False)

    def test_invalid_dates_counts_durations_and_dictionary_values_fail(self) -> None:
        base = [
            "batch-1", "UR3", "real_robot_teleoperation", "single_arm", "liquid_pour",
            "液体倾倒", "pour_water", "description", date(2026, 4, 29), 10,
            timedelta(seconds=20), "operator01", "",
        ]
        cases = (
            (8, "2026-02-30", "calendar date"),
            (9, -1, "positive integer"),
            (10, "00:60:00", r"\[h\]:mm:ss"),
            (1, "NOT_A_DEVICE", "unknown or disabled device"),
            (5, "错误任务名", "must equal"),
            (12, "#VALUE!", "contains Excel error"),
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            for index, (column, invalid_value, message) in enumerate(cases):
                with self.subTest(column=column, value=invalid_value):
                    row = [*base]
                    row[column] = invalid_value
                    workbook_path = Path(temp_dir) / f"invalid-{index}.xlsx"
                    _make_workbook(workbook_path, [row], [])
                    with self.assertRaisesRegex(DatasetStatisticsValidationError, message):
                        build_statistics_payload(workbook_path, validate_summary=False)

    def test_duplicate_simulation_dataset_id_fails(self) -> None:
        simulation = [
            [
                "sim-batch-1", "dataset-1", "ISAAC_SIM", "FRANKA", "liquid_pour",
                "液体倾倒", "scene-a", date(2026, 7, 1), 10, 1000,
                timedelta(seconds=100), "",
            ],
            [
                "sim-batch-2", "dataset-1", "ISAAC_SIM", "FRANKA", "liquid_pour",
                "液体倾倒", "scene-b", date(2026, 7, 2), 10, 1000,
                timedelta(seconds=100), "",
            ],
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "duplicate-dataset.xlsx"
            _make_workbook(workbook_path, [], simulation)
            with self.assertRaisesRegex(DatasetStatisticsValidationError, "dataset_id.*duplicated"):
                build_statistics_payload(workbook_path, validate_summary=False)

    def test_xarm_is_emitted_only_after_a_nonzero_detail_row_exists(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            empty_path = Path(temp_dir) / "empty.xlsx"
            _make_workbook(empty_path, [], [])
            empty_payload = build_statistics_payload(empty_path, validate_summary=False)
            self.assertNotIn(
                "XARM",
                {row["device_model"] for row in empty_payload["real_robot_teleoperation"]["rows"]},
            )
            xarm_path = Path(temp_dir) / "xarm.xlsx"
            _make_workbook(
                xarm_path,
                [[
                    "xarm-batch-1", "XARM", "real_robot_teleoperation", "single_arm",
                    "liquid_pour", "液体倾倒", "pour_water", "description",
                    date(2026, 7, 17), 12, timedelta(seconds=30), "operator01", "",
                ]],
                [],
            )
            xarm_payload = build_statistics_payload(xarm_path, validate_summary=False)
            self.assertEqual(
                [row["device_model"] for row in xarm_payload["real_robot_teleoperation"]["rows"]],
                ["XARM"],
            )

    def test_json_write_is_atomic_semantic_noop_and_check_detects_drift(self) -> None:
        rows = [
            [
                "batch-1", "UR3", "real_robot_teleoperation", "single_arm", "liquid_pour",
                "液体倾倒", "pour_water", "description", date(2026, 4, 29), 10,
                timedelta(seconds=20), "operator01", "",
            ]
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "registry.xlsx"
            output_path = Path(temp_dir) / "statistics.json"
            _make_workbook(workbook_path, rows, [])
            first = generate_statistics_json(
                workbook_path, output_path, validate_summary=False
            )
            original_bytes = output_path.read_bytes()
            second = generate_statistics_json(
                workbook_path, output_path, validate_summary=False
            )
            self.assertTrue(first.changed)
            self.assertFalse(second.changed)
            self.assertEqual(output_path.read_bytes(), original_bytes)
            payload = json.loads(output_path.read_text(encoding="utf-8"))
            payload["real_robot_teleoperation"]["totals"]["sample_count"] = 999
            output_path.write_text(json.dumps(payload), encoding="utf-8")
            with self.assertRaisesRegex(DatasetStatisticsValidationError, "out of sync"):
                generate_statistics_json(
                    workbook_path, output_path, check=True, validate_summary=False
                )

            invalid_workbook = Path(temp_dir) / "invalid-registry.xlsx"
            invalid_row = [*rows[0]]
            invalid_row[9] = -1
            _make_workbook(invalid_workbook, [invalid_row], [])
            before_invalid_build = output_path.read_bytes()
            with self.assertRaises(DatasetStatisticsValidationError):
                generate_statistics_json(
                    invalid_workbook, output_path, validate_summary=False
                )
            self.assertEqual(output_path.read_bytes(), before_invalid_build)

    def test_committed_registry_matches_expected_baseline(self) -> None:
        if not DEFAULT_WORKBOOK_PATH.is_file():
            self.skipTest("registry workbook has not been generated yet")
        payload = build_statistics_payload(DEFAULT_WORKBOOK_PATH)
        real = payload["real_robot_teleoperation"]
        umi = payload["umi"]
        self.assertEqual(real["totals"], {"batch_count": 15, "sample_count": 873, "duration_seconds": 17468})
        self.assertEqual(umi["totals"], {"batch_count": 12, "sample_count": 743, "duration_seconds": 9322})
        self.assertEqual(real["totals"]["batch_count"] + umi["totals"]["batch_count"], 27)
        self.assertEqual(real["totals"]["sample_count"] + umi["totals"]["sample_count"], 1616)
        self.assertEqual(real["totals"]["duration_seconds"] + umi["totals"]["duration_seconds"], 26790)
        self.assertNotIn("XARM", {row["device_model"] for row in real["rows"]})
        centrifuge = next(
            row for row in real["rows"]
            if row["device_model"] == "UR3" and row["task_id"] == "centrifuge_tube_to_rack"
        )
        self.assertEqual(centrifuge["batch_count"], 3)
        self.assertEqual(centrifuge["sample_count"], 260)
        self.assertEqual(centrifuge["duration_seconds"], 4555)
        self.assertEqual(centrifuge["collection_date_start"], "2026-06-23")
        self.assertEqual(centrifuge["collection_date_end"], "2026-07-06")


if __name__ == "__main__":
    unittest.main()
