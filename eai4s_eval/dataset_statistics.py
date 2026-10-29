"""Validation and aggregation for the dataset statistics registry workbook.

The workbook is the editable source of truth.  This module deliberately keeps
the generated JSON small and presentation-neutral so that the Flask dashboard
can reload it without importing openpyxl on every request.
"""

from __future__ import annotations

import json
import math
import os
import re
import tempfile
from collections import OrderedDict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence


SCHEMA_VERSION = "dataset_statistics/v1"
PACKAGE_ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK_PATH = PACKAGE_ROOT / "data" / "dataset_statistics_registry.xlsx"
DEFAULT_JSON_PATH = PACKAGE_ROOT / "data" / "dataset_statistics.json"

INSTRUCTIONS_SHEET = "使用说明"
PHYSICAL_SHEET = "真机UMI明细"
SIMULATION_SHEET = "仿真明细"
DICTIONARY_SHEET = "配置字典"
SUMMARY_SHEET = "统计汇总"
REQUIRED_SHEETS = (
    INSTRUCTIONS_SHEET,
    PHYSICAL_SHEET,
    SIMULATION_SHEET,
    DICTIONARY_SHEET,
    SUMMARY_SHEET,
)

PHYSICAL_HEADERS = (
    "batch_id",
    "device_model",
    "collection_method",
    "arm_mode",
    "task_id",
    "task_name",
    "raw_task_name",
    "task_description",
    "collection_date",
    "sample_count",
    "valid_duration",
    "operator_ids",
    "notes",
)
SIMULATION_HEADERS = (
    "batch_id",
    "dataset_id",
    "simulator",
    "robot_type",
    "task_id",
    "task_name",
    "scene_id",
    "generation_date",
    "episode_count",
    "frame_count",
    "simulation_duration",
    "notes",
)
DICTIONARY_HEADERS = (
    "dictionary_type",
    "code",
    "name",
    "collection_method",
    "arm_mode",
    "enabled",
    "notes",
)

REQUIRED_TASKS = OrderedDict(
    (
        ("liquid_pour", "液体倾倒"),
        ("test_tube_to_rack", "试管入架"),
        ("centrifuge_tube_to_rack", "离心管入架"),
        ("sample_vial_storage", "样品瓶收纳"),
        ("glass_beaker_pick_place", "玻璃烧杯抓放"),
        ("clear_bottle_transfer", "透明瓶转移"),
    )
)
REQUIRED_DEVICE_METHODS = {
    "XARM": "real_robot_teleoperation",
    "UR3": "real_robot_teleoperation",
    "OPENARM": "real_robot_teleoperation",
    "UMI": "umi",
}
OPERATOR_RE = re.compile(r"^operator\d{2}$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
DURATION_RE = re.compile(r"^(\d+):([0-5]\d):([0-5]\d)$")


class DatasetStatisticsValidationError(ValueError):
    """Raised when the registry, its formulas, or generated JSON is invalid."""


@dataclass(frozen=True)
class BuildResult:
    payload: dict[str, Any]
    changed: bool


@dataclass(frozen=True)
class _DictionaryEntry:
    dictionary_type: str
    code: str
    name: str
    collection_method: str
    arm_mode: str
    enabled: bool
    row_number: int


def _require_openpyxl() -> tuple[Any, Any]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.datetime import from_excel
    except ImportError as exc:  # pragma: no cover - depends on runtime setup
        raise RuntimeError(
            "Reading dataset_statistics_registry.xlsx requires openpyxl. "
            "Install the dashboard requirements first."
        ) from exc
    return load_workbook, from_excel


def _clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _cell_context(sheet_name: str, row_number: int, field: str) -> str:
    return f"{sheet_name}!{field} (row {row_number})"


def _required_text(value: Any, *, sheet: str, row: int, field: str) -> str:
    text = _clean_text(value)
    if not text:
        raise DatasetStatisticsValidationError(
            f"{_cell_context(sheet, row, field)} must not be empty"
        )
    if text.startswith("="):
        raise DatasetStatisticsValidationError(
            f"{_cell_context(sheet, row, field)} must be a value, not a formula"
        )
    return text


def _parse_enabled(value: Any, *, sheet: str, row: int) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value in (0, 1):
            return bool(value)
    text = _clean_text(value).lower()
    if text in {"1", "true", "yes", "y", "是", "启用", "enabled"}:
        return True
    if text in {"0", "false", "no", "n", "否", "停用", "disabled"}:
        return False
    raise DatasetStatisticsValidationError(
        f"{_cell_context(sheet, row, 'enabled')} must be a boolean/yes/no value"
    )


def _parse_date(value: Any, *, sheet: str, row: int, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        _, from_excel = _require_openpyxl()
        try:
            converted = from_excel(value)
        except (TypeError, ValueError, OverflowError) as exc:
            raise DatasetStatisticsValidationError(
                f"{_cell_context(sheet, row, field)} is not a valid Excel date"
            ) from exc
        if isinstance(converted, datetime):
            return converted.date()
        if isinstance(converted, date):
            return converted
    text = _clean_text(value)
    if not DATE_RE.fullmatch(text):
        raise DatasetStatisticsValidationError(
            f"{_cell_context(sheet, row, field)} must use YYYY-MM-DD"
        )
    try:
        return date.fromisoformat(text)
    except ValueError as exc:
        raise DatasetStatisticsValidationError(
            f"{_cell_context(sheet, row, field)} is not a valid calendar date"
        ) from exc


def _parse_nonnegative_int(
    value: Any,
    *,
    sheet: str,
    row: int,
    field: str,
    positive: bool = False,
) -> int:
    if isinstance(value, bool):
        parsed: int | None = None
    elif isinstance(value, int):
        parsed = value
    elif isinstance(value, float) and math.isfinite(value) and value.is_integer():
        parsed = int(value)
    else:
        text = _clean_text(value)
        parsed = int(text) if re.fullmatch(r"\d+", text) else None
    lower_bound = 1 if positive else 0
    if parsed is None or parsed < lower_bound:
        qualifier = "a positive" if positive else "a non-negative"
        raise DatasetStatisticsValidationError(
            f"{_cell_context(sheet, row, field)} must be {qualifier} integer"
        )
    return parsed


def duration_to_seconds(
    value: Any,
    *,
    sheet: str = "workbook",
    row: int = 0,
    field: str = "duration",
) -> int:
    """Convert an Excel duration or ``H:MM:SS`` string to whole seconds."""

    seconds: float | int
    if isinstance(value, timedelta):
        seconds = value.total_seconds()
    elif isinstance(value, time):
        seconds = value.hour * 3600 + value.minute * 60 + value.second
        seconds += value.microsecond / 1_000_000
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(float(value)):
            seconds = -1
        else:
            # Excel stores durations as fractions of a 24-hour day.
            seconds = float(value) * 86400
    else:
        text = _clean_text(value)
        match = DURATION_RE.fullmatch(text)
        if not match:
            raise DatasetStatisticsValidationError(
                f"{_cell_context(sheet, row, field)} must use [h]:mm:ss"
            )
        hours, minutes, secs = (int(part) for part in match.groups())
        seconds = hours * 3600 + minutes * 60 + secs
    if seconds < 0 or abs(seconds - round(seconds)) > 1e-6:
        raise DatasetStatisticsValidationError(
            f"{_cell_context(sheet, row, field)} must be a non-negative whole-second duration"
        )
    return int(round(seconds))


def _parse_operators(value: Any, *, sheet: str, row: int) -> list[str]:
    text = _required_text(value, sheet=sheet, row=row, field="operator_ids")
    operators = [item.strip() for item in text.split(";")]
    if any(not OPERATOR_RE.fullmatch(item) for item in operators):
        raise DatasetStatisticsValidationError(
            f"{_cell_context(sheet, row, 'operator_ids')} must contain semicolon-separated operatorNN ids"
        )
    if len(set(operators)) != len(operators):
        raise DatasetStatisticsValidationError(
            f"{_cell_context(sheet, row, 'operator_ids')} contains duplicate ids"
        )
    return operators


def _find_header_row(ws: Any, required_headers: Sequence[str]) -> tuple[int, dict[str, Any]]:
    required = set(required_headers)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        cells = {_clean_text(cell.value): cell for cell in row if _clean_text(cell.value)}
        if required.issubset(cells):
            return row[0].row, {name: cells[name] for name in required_headers}
    raise DatasetStatisticsValidationError(
        f"{ws.title} is missing the required header row: {', '.join(required_headers)}"
    )


def _read_table(ws: Any, required_headers: Sequence[str]) -> list[dict[str, Any]]:
    header_row, header_cells = _find_header_row(ws, required_headers)
    columns = {name: cell.column for name, cell in header_cells.items()}
    records: list[dict[str, Any]] = []
    for row_number in range(header_row + 1, ws.max_row + 1):
        cells = {name: ws.cell(row=row_number, column=column) for name, column in columns.items()}
        if all(cell.value in (None, "") for cell in cells.values()):
            continue
        record = {name: cell.value for name, cell in cells.items()}
        record["__row__"] = row_number
        record["__cells__"] = cells
        records.append(record)
    return records


def _validate_no_excel_errors(workbooks: Sequence[Any]) -> None:
    seen: set[tuple[str, str]] = set()
    for workbook in workbooks:
        for sheet_name in REQUIRED_SHEETS:
            ws = workbook[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type != "e":
                        continue
                    key = (sheet_name, cell.coordinate)
                    if key in seen:
                        continue
                    seen.add(key)
                    raise DatasetStatisticsValidationError(
                        f"{sheet_name}!{cell.coordinate} contains Excel error {cell.value!r}"
                    )


def _load_dictionaries(ws: Any) -> dict[str, OrderedDict[str, _DictionaryEntry]]:
    registry: dict[str, OrderedDict[str, _DictionaryEntry]] = {}
    for record in _read_table(ws, DICTIONARY_HEADERS):
        row = int(record["__row__"])
        dictionary_type = _required_text(
            record["dictionary_type"], sheet=DICTIONARY_SHEET, row=row, field="dictionary_type"
        ).lower()
        code = _required_text(record["code"], sheet=DICTIONARY_SHEET, row=row, field="code")
        name = _required_text(record["name"], sheet=DICTIONARY_SHEET, row=row, field="name")
        collection_method = _clean_text(record["collection_method"])
        arm_mode = _clean_text(record["arm_mode"])
        enabled = _parse_enabled(record["enabled"], sheet=DICTIONARY_SHEET, row=row)
        entries = registry.setdefault(dictionary_type, OrderedDict())
        if code in entries:
            raise DatasetStatisticsValidationError(
                f"{DICTIONARY_SHEET}!code {code!r} is duplicated for dictionary_type {dictionary_type!r}"
            )
        entries[code] = _DictionaryEntry(
            dictionary_type=dictionary_type,
            code=code,
            name=name,
            collection_method=collection_method,
            arm_mode=arm_mode,
            enabled=enabled,
            row_number=row,
        )

    for dictionary_type in ("task", "device", "simulator", "robot"):
        if not registry.get(dictionary_type):
            raise DatasetStatisticsValidationError(
                f"{DICTIONARY_SHEET} must define at least one {dictionary_type!r} entry"
            )
    tasks = registry["task"]
    for task_id, expected_name in REQUIRED_TASKS.items():
        entry = tasks.get(task_id)
        if entry is None or entry.name != expected_name or not entry.enabled:
            raise DatasetStatisticsValidationError(
                f"{DICTIONARY_SHEET} must enable task {task_id!r} with name {expected_name!r}"
            )
    devices = registry["device"]
    for device, expected_method in REQUIRED_DEVICE_METHODS.items():
        entry = devices.get(device)
        if entry is None or entry.collection_method != expected_method or not entry.enabled:
            raise DatasetStatisticsValidationError(
                f"{DICTIONARY_SHEET} must enable device {device!r} with collection_method {expected_method!r}"
            )
        if not entry.arm_mode:
            raise DatasetStatisticsValidationError(
                f"{DICTIONARY_SHEET} device {device!r} must define arm_mode"
            )
    return registry


def _enabled_entry(
    dictionaries: Mapping[str, Mapping[str, _DictionaryEntry]],
    dictionary_type: str,
    code: str,
    *,
    sheet: str,
    row: int,
    field: str,
) -> _DictionaryEntry:
    entry = dictionaries.get(dictionary_type, {}).get(code)
    if entry is None or not entry.enabled:
        raise DatasetStatisticsValidationError(
            f"{_cell_context(sheet, row, field)} references unknown or disabled {dictionary_type} {code!r}"
        )
    return entry


def _validate_physical_rows(
    ws: Any, dictionaries: Mapping[str, Mapping[str, _DictionaryEntry]]
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for record in _read_table(ws, PHYSICAL_HEADERS):
        row = int(record["__row__"])
        batch_id = _required_text(record["batch_id"], sheet=PHYSICAL_SHEET, row=row, field="batch_id")
        device = _required_text(record["device_model"], sheet=PHYSICAL_SHEET, row=row, field="device_model")
        device_entry = _enabled_entry(
            dictionaries, "device", device, sheet=PHYSICAL_SHEET, row=row, field="device_model"
        )
        method = _required_text(
            record["collection_method"], sheet=PHYSICAL_SHEET, row=row, field="collection_method"
        )
        if method != device_entry.collection_method or method not in {"real_robot_teleoperation", "umi"}:
            raise DatasetStatisticsValidationError(
                f"{_cell_context(PHYSICAL_SHEET, row, 'collection_method')} must equal the device dictionary value"
            )
        arm_mode = _required_text(record["arm_mode"], sheet=PHYSICAL_SHEET, row=row, field="arm_mode")
        if arm_mode != device_entry.arm_mode:
            raise DatasetStatisticsValidationError(
                f"{_cell_context(PHYSICAL_SHEET, row, 'arm_mode')} must equal the device dictionary value"
            )
        task_id = _required_text(record["task_id"], sheet=PHYSICAL_SHEET, row=row, field="task_id")
        task_entry = _enabled_entry(
            dictionaries, "task", task_id, sheet=PHYSICAL_SHEET, row=row, field="task_id"
        )
        task_name = _required_text(record["task_name"], sheet=PHYSICAL_SHEET, row=row, field="task_name")
        if task_name != task_entry.name:
            raise DatasetStatisticsValidationError(
                f"{_cell_context(PHYSICAL_SHEET, row, 'task_name')} must equal {task_entry.name!r}"
            )
        raw_task_name = _required_text(
            record["raw_task_name"], sheet=PHYSICAL_SHEET, row=row, field="raw_task_name"
        )
        task_description = _required_text(
            record["task_description"], sheet=PHYSICAL_SHEET, row=row, field="task_description"
        )
        collection_date = _parse_date(
            record["collection_date"], sheet=PHYSICAL_SHEET, row=row, field="collection_date"
        )
        sample_count = _parse_nonnegative_int(
            record["sample_count"], sheet=PHYSICAL_SHEET, row=row, field="sample_count", positive=True
        )
        duration_seconds = duration_to_seconds(
            record["valid_duration"], sheet=PHYSICAL_SHEET, row=row, field="valid_duration"
        )
        operators = _parse_operators(record["operator_ids"], sheet=PHYSICAL_SHEET, row=row)
        result.append(
            {
                "batch_id": batch_id,
                "device_model": device,
                "collection_method": method,
                "arm_mode": arm_mode,
                "task_id": task_id,
                "task_name": task_name,
                "raw_task_name": raw_task_name,
                "task_description": task_description,
                "collection_date": collection_date,
                "sample_count": sample_count,
                "duration_seconds": duration_seconds,
                "operator_ids": operators,
                "notes": _clean_text(record["notes"]),
                "__row__": row,
            }
        )
    return result


def _validate_simulation_rows(
    ws: Any, dictionaries: Mapping[str, Mapping[str, _DictionaryEntry]]
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for record in _read_table(ws, SIMULATION_HEADERS):
        row = int(record["__row__"])
        batch_id = _required_text(record["batch_id"], sheet=SIMULATION_SHEET, row=row, field="batch_id")
        dataset_id = _required_text(record["dataset_id"], sheet=SIMULATION_SHEET, row=row, field="dataset_id")
        simulator = _required_text(record["simulator"], sheet=SIMULATION_SHEET, row=row, field="simulator")
        _enabled_entry(
            dictionaries, "simulator", simulator, sheet=SIMULATION_SHEET, row=row, field="simulator"
        )
        robot_type = _required_text(record["robot_type"], sheet=SIMULATION_SHEET, row=row, field="robot_type")
        _enabled_entry(
            dictionaries, "robot", robot_type, sheet=SIMULATION_SHEET, row=row, field="robot_type"
        )
        task_id = _required_text(record["task_id"], sheet=SIMULATION_SHEET, row=row, field="task_id")
        task_entry = _enabled_entry(
            dictionaries, "task", task_id, sheet=SIMULATION_SHEET, row=row, field="task_id"
        )
        task_name = _required_text(record["task_name"], sheet=SIMULATION_SHEET, row=row, field="task_name")
        if task_name != task_entry.name:
            raise DatasetStatisticsValidationError(
                f"{_cell_context(SIMULATION_SHEET, row, 'task_name')} must equal {task_entry.name!r}"
            )
        scene_id = _required_text(record["scene_id"], sheet=SIMULATION_SHEET, row=row, field="scene_id")
        generation_date = _parse_date(
            record["generation_date"], sheet=SIMULATION_SHEET, row=row, field="generation_date"
        )
        episode_count = _parse_nonnegative_int(
            record["episode_count"], sheet=SIMULATION_SHEET, row=row, field="episode_count"
        )
        frame_count = _parse_nonnegative_int(
            record["frame_count"], sheet=SIMULATION_SHEET, row=row, field="frame_count"
        )
        duration_seconds = duration_to_seconds(
            record["simulation_duration"], sheet=SIMULATION_SHEET, row=row, field="simulation_duration"
        )
        result.append(
            {
                "batch_id": batch_id,
                "dataset_id": dataset_id,
                "simulator": simulator,
                "robot_type": robot_type,
                "task_id": task_id,
                "task_name": task_name,
                "scene_id": scene_id,
                "generation_date": generation_date,
                "episode_count": episode_count,
                "frame_count": frame_count,
                "duration_seconds": duration_seconds,
                "notes": _clean_text(record["notes"]),
                "__row__": row,
            }
        )
    return result


def _validate_unique_ids(physical_rows: Sequence[Mapping[str, Any]], simulation_rows: Sequence[Mapping[str, Any]]) -> None:
    batch_ids: dict[str, tuple[str, int]] = {}
    for sheet, rows in ((PHYSICAL_SHEET, physical_rows), (SIMULATION_SHEET, simulation_rows)):
        for row in rows:
            batch_id = str(row["batch_id"])
            location = (sheet, int(row["__row__"]))
            if batch_id in batch_ids:
                previous_sheet, previous_row = batch_ids[batch_id]
                raise DatasetStatisticsValidationError(
                    f"batch_id {batch_id!r} is duplicated at {previous_sheet} row {previous_row} and {sheet} row {location[1]}"
                )
            batch_ids[batch_id] = location
    dataset_ids: dict[str, int] = {}
    for row in simulation_rows:
        dataset_id = str(row["dataset_id"])
        row_number = int(row["__row__"])
        if dataset_id in dataset_ids:
            raise DatasetStatisticsValidationError(
                f"dataset_id {dataset_id!r} is duplicated at {SIMULATION_SHEET} rows {dataset_ids[dataset_id]} and {row_number}"
            )
        dataset_ids[dataset_id] = row_number


def _aggregate_physical(rows: Iterable[Mapping[str, Any]], method: str) -> dict[str, Any]:
    groups: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()
    for row in rows:
        if row["collection_method"] != method:
            continue
        key = (str(row["device_model"]), str(row["task_id"]))
        group = groups.get(key)
        if group is None:
            group = {
                "device_model": key[0],
                "task_id": key[1],
                "task_name": str(row["task_name"]),
                "batch_count": 0,
                "sample_count": 0,
                "duration_seconds": 0,
                "collection_date_start": row["collection_date"],
                "collection_date_end": row["collection_date"],
                "operator_ids": [],
            }
            groups[key] = group
        group["batch_count"] += 1
        group["sample_count"] += int(row["sample_count"])
        group["duration_seconds"] += int(row["duration_seconds"])
        group["collection_date_start"] = min(group["collection_date_start"], row["collection_date"])
        group["collection_date_end"] = max(group["collection_date_end"], row["collection_date"])
        for operator in row["operator_ids"]:
            if operator not in group["operator_ids"]:
                group["operator_ids"].append(operator)
    output_rows: list[dict[str, Any]] = []
    for group in groups.values():
        group["collection_date_start"] = group["collection_date_start"].isoformat()
        group["collection_date_end"] = group["collection_date_end"].isoformat()
        output_rows.append(group)
    totals = {
        "batch_count": sum(row["batch_count"] for row in output_rows),
        "sample_count": sum(row["sample_count"] for row in output_rows),
        "duration_seconds": sum(row["duration_seconds"] for row in output_rows),
    }
    return {"status": "active" if output_rows else "empty", "totals": totals, "rows": output_rows}


def _aggregate_simulation(rows: Iterable[Mapping[str, Any]]) -> dict[str, Any]:
    source_rows = list(rows)
    groups: OrderedDict[tuple[str, str, str], dict[str, Any]] = OrderedDict()
    for row in source_rows:
        key = (str(row["simulator"]), str(row["robot_type"]), str(row["task_id"]))
        group = groups.get(key)
        if group is None:
            group = {
                "simulator": key[0],
                "robot_type": key[1],
                "task_id": key[2],
                "task_name": str(row["task_name"]),
                "dataset_batch_count": 0,
                "episode_count": 0,
                "frame_count": 0,
                "duration_seconds": 0,
                "scene_count": 0,
                "generation_date_start": row["generation_date"],
                "generation_date_end": row["generation_date"],
                "_scene_ids": [],
            }
            groups[key] = group
        group["dataset_batch_count"] += 1
        group["episode_count"] += int(row["episode_count"])
        group["frame_count"] += int(row["frame_count"])
        group["duration_seconds"] += int(row["duration_seconds"])
        group["generation_date_start"] = min(group["generation_date_start"], row["generation_date"])
        group["generation_date_end"] = max(group["generation_date_end"], row["generation_date"])
        if row["scene_id"] not in group["_scene_ids"]:
            group["_scene_ids"].append(row["scene_id"])
    output_rows: list[dict[str, Any]] = []
    for group in groups.values():
        group["scene_count"] = len(group.pop("_scene_ids"))
        group["generation_date_start"] = group["generation_date_start"].isoformat()
        group["generation_date_end"] = group["generation_date_end"].isoformat()
        output_rows.append(group)
    totals = {
        "dataset_batch_count": len(source_rows),
        "episode_count": sum(int(row["episode_count"]) for row in source_rows),
        "frame_count": sum(int(row["frame_count"]) for row in source_rows),
        "duration_seconds": sum(int(row["duration_seconds"]) for row in source_rows),
        "scene_count": len({str(row["scene_id"]) for row in source_rows}),
    }
    return {"status": "active" if output_rows else "empty", "totals": totals, "rows": output_rows}


def _build_sections(physical_rows: Sequence[Mapping[str, Any]], simulation_rows: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    return {
        "real_robot_teleoperation": _aggregate_physical(physical_rows, "real_robot_teleoperation"),
        "umi": _aggregate_physical(physical_rows, "umi"),
        "simulation": _aggregate_simulation(simulation_rows),
    }


def _marker_key(value: Any) -> str:
    return re.sub(r"\s+", "", _clean_text(value)).lower()


SUMMARY_MARKERS = {
    "真机遥操作汇总": "real_robot_teleoperation",
    "umi汇总": "umi",
    "仿真汇总": "simulation",
}


def _summary_sections(ws: Any) -> dict[str, tuple[int, int]]:
    markers: list[tuple[int, str]] = []
    for row in ws.iter_rows():
        for cell in row:
            section = SUMMARY_MARKERS.get(_marker_key(cell.value))
            if section is not None:
                markers.append((cell.row, section))
                break
    found = {section for _, section in markers}
    if found != set(SUMMARY_MARKERS.values()):
        missing = sorted(set(SUMMARY_MARKERS.values()) - found)
        raise DatasetStatisticsValidationError(
            f"{SUMMARY_SHEET} is missing summary section marker(s): {', '.join(missing)}"
        )
    markers.sort()
    bounds: dict[str, tuple[int, int]] = {}
    for index, (start, section) in enumerate(markers):
        end = markers[index + 1][0] - 1 if index + 1 < len(markers) else ws.max_row
        bounds[section] = (start, end)
    return bounds


def _summary_header(ws: Any, start_row: int, end_row: int, expected: Sequence[str]) -> tuple[int, dict[str, int]]:
    expected_set = set(expected)
    for row_number in range(start_row + 1, min(end_row, start_row + 5) + 1):
        values = {_clean_text(cell.value): cell.column for cell in ws[row_number] if _clean_text(cell.value)}
        if expected_set.issubset(values):
            return row_number, {name: values[name] for name in expected}
    raise DatasetStatisticsValidationError(
        f"{SUMMARY_SHEET} section at row {start_row} is missing headers: {', '.join(expected)}"
    )


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _assert_summary_formula_cells(
    formula_ws: Any,
    row_number: int,
    columns: Mapping[str, int],
    fields: Sequence[str],
) -> None:
    missing = [field for field in fields if not _is_formula(formula_ws.cell(row_number, columns[field]).value)]
    if missing:
        raise DatasetStatisticsValidationError(
            f"{SUMMARY_SHEET} row {row_number} must use formulas for: {', '.join(missing)}"
        )


def _summary_cached_value(cached_ws: Any, row: int, column: int, field: str) -> Any:
    value = cached_ws.cell(row, column).value
    if value is None:
        raise DatasetStatisticsValidationError(
            f"{SUMMARY_SHEET}!{field} (row {row}) has no cached formula value; recalculate and save the workbook"
        )
    return value


def _compare_value(actual: Any, expected: Any, *, row: int, field: str) -> None:
    if actual != expected:
        raise DatasetStatisticsValidationError(
            f"{SUMMARY_SHEET}!{field} (row {row}) is {actual!r}, expected {expected!r}"
        )


def _reconcile_physical_summary(
    formula_ws: Any,
    cached_ws: Any,
    section_name: str,
    section: Mapping[str, Any],
    start_row: int,
    end_row: int,
) -> None:
    if section_name == "real_robot_teleoperation":
        headers = (
            "device_model", "task_id", "task_name", "batch_count", "sample_count",
            "duration", "first_date", "last_date", "operator_ids",
        )
        key_fields = ("device_model", "task_id")
        total_label = "真机遥操作总计"
    else:
        headers = (
            "task_id", "task_name", "batch_count", "sample_count", "duration",
            "first_date", "last_date", "operator_ids",
        )
        key_fields = ("task_id",)
        total_label = "UMI总计"
    header_row, columns = _summary_header(formula_ws, start_row, end_row, headers)
    expected_rows = {
        tuple(str(row[field]) for field in key_fields): row for row in section["rows"]
    }
    seen: set[tuple[str, ...]] = set()
    total_row: int | None = None
    for row_number in range(header_row + 1, end_row + 1):
        first_value = _clean_text(formula_ws.cell(row_number, columns[key_fields[0]]).value)
        if _marker_key(first_value) == _marker_key(total_label):
            total_row = row_number
            break
        if not first_value:
            continue
        key = tuple(_clean_text(formula_ws.cell(row_number, columns[field]).value) for field in key_fields)
        expected = expected_rows.get(key)
        if expected is None:
            raise DatasetStatisticsValidationError(
                f"{SUMMARY_SHEET} row {row_number} contains unexpected group {key!r}"
            )
        if key in seen:
            raise DatasetStatisticsValidationError(f"{SUMMARY_SHEET} repeats group {key!r}")
        seen.add(key)
        _assert_summary_formula_cells(
            formula_ws, row_number, columns,
            ("batch_count", "sample_count", "duration", "first_date", "last_date"),
        )
        task_name = _clean_text(_summary_cached_value(cached_ws, row_number, columns["task_name"], "task_name"))
        batch_count = _parse_nonnegative_int(
            _summary_cached_value(cached_ws, row_number, columns["batch_count"], "batch_count"),
            sheet=SUMMARY_SHEET, row=row_number, field="batch_count",
        )
        sample_count = _parse_nonnegative_int(
            _summary_cached_value(cached_ws, row_number, columns["sample_count"], "sample_count"),
            sheet=SUMMARY_SHEET, row=row_number, field="sample_count",
        )
        duration = duration_to_seconds(
            _summary_cached_value(cached_ws, row_number, columns["duration"], "duration"),
            sheet=SUMMARY_SHEET, row=row_number, field="duration",
        )
        first_date = _parse_date(
            _summary_cached_value(cached_ws, row_number, columns["first_date"], "first_date"),
            sheet=SUMMARY_SHEET, row=row_number, field="first_date",
        ).isoformat()
        last_date = _parse_date(
            _summary_cached_value(cached_ws, row_number, columns["last_date"], "last_date"),
            sheet=SUMMARY_SHEET, row=row_number, field="last_date",
        ).isoformat()
        operators = _parse_operators(
            _summary_cached_value(cached_ws, row_number, columns["operator_ids"], "operator_ids"),
            sheet=SUMMARY_SHEET, row=row_number,
        )
        _compare_value(task_name, expected["task_name"], row=row_number, field="task_name")
        _compare_value(batch_count, expected["batch_count"], row=row_number, field="batch_count")
        _compare_value(sample_count, expected["sample_count"], row=row_number, field="sample_count")
        _compare_value(duration, expected["duration_seconds"], row=row_number, field="duration")
        _compare_value(first_date, expected["collection_date_start"], row=row_number, field="first_date")
        _compare_value(last_date, expected["collection_date_end"], row=row_number, field="last_date")
        _compare_value(operators, expected["operator_ids"], row=row_number, field="operator_ids")
    missing_groups = sorted(set(expected_rows) - seen)
    if missing_groups:
        raise DatasetStatisticsValidationError(
            f"{SUMMARY_SHEET} is missing {section_name} group row(s): {missing_groups!r}; "
            "add the corresponding formula rows before rebuilding JSON"
        )
    if total_row is None:
        raise DatasetStatisticsValidationError(f"{SUMMARY_SHEET} is missing {total_label}")
    _assert_summary_formula_cells(
        formula_ws,
        total_row,
        columns,
        ("batch_count", "sample_count", "duration", "first_date", "last_date"),
    )
    totals = section["totals"]
    for field, expected_field in (
        ("batch_count", "batch_count"),
        ("sample_count", "sample_count"),
    ):
        actual = _parse_nonnegative_int(
            _summary_cached_value(cached_ws, total_row, columns[field], field),
            sheet=SUMMARY_SHEET, row=total_row, field=field,
        )
        _compare_value(actual, totals[expected_field], row=total_row, field=field)
    duration = duration_to_seconds(
        _summary_cached_value(cached_ws, total_row, columns["duration"], "duration"),
        sheet=SUMMARY_SHEET, row=total_row, field="duration",
    )
    _compare_value(duration, totals["duration_seconds"], row=total_row, field="duration")
    if section["rows"]:
        first_date = _parse_date(
            _summary_cached_value(cached_ws, total_row, columns["first_date"], "first_date"),
            sheet=SUMMARY_SHEET,
            row=total_row,
            field="first_date",
        ).isoformat()
        last_date = _parse_date(
            _summary_cached_value(cached_ws, total_row, columns["last_date"], "last_date"),
            sheet=SUMMARY_SHEET,
            row=total_row,
            field="last_date",
        ).isoformat()
        expected_first = min(str(item["collection_date_start"]) for item in section["rows"])
        expected_last = max(str(item["collection_date_end"]) for item in section["rows"])
        _compare_value(first_date, expected_first, row=total_row, field="first_date")
        _compare_value(last_date, expected_last, row=total_row, field="last_date")
        expected_operators: list[str] = []
        for item in section["rows"]:
            for operator in item["operator_ids"]:
                if operator not in expected_operators:
                    expected_operators.append(operator)
        total_operators = _parse_operators(
            _summary_cached_value(cached_ws, total_row, columns["operator_ids"], "operator_ids"),
            sheet=SUMMARY_SHEET,
            row=total_row,
        )
        _compare_value(total_operators, expected_operators, row=total_row, field="operator_ids")


def _reconcile_simulation_summary(
    formula_ws: Any,
    cached_ws: Any,
    section: Mapping[str, Any],
    start_row: int,
    end_row: int,
) -> None:
    headers = (
        "simulator", "robot_type", "task_id", "task_name", "dataset_batch_count",
        "episode_count", "frame_count", "duration", "scene_count", "first_date", "last_date",
    )
    header_row, columns = _summary_header(formula_ws, start_row, end_row, headers)
    expected_rows = {
        (str(row["simulator"]), str(row["robot_type"]), str(row["task_id"])): row
        for row in section["rows"]
    }
    seen: set[tuple[str, str, str]] = set()
    total_row: int | None = None
    for row_number in range(header_row + 1, end_row + 1):
        simulator = _clean_text(formula_ws.cell(row_number, columns["simulator"]).value)
        if _marker_key(simulator) == _marker_key("仿真总计"):
            total_row = row_number
            break
        if not simulator or "暂无仿真数据" in simulator:
            continue
        key = (
            simulator,
            _clean_text(formula_ws.cell(row_number, columns["robot_type"]).value),
            _clean_text(formula_ws.cell(row_number, columns["task_id"]).value),
        )
        expected = expected_rows.get(key)
        if expected is None:
            raise DatasetStatisticsValidationError(
                f"{SUMMARY_SHEET} row {row_number} contains unexpected simulation group {key!r}"
            )
        if key in seen:
            raise DatasetStatisticsValidationError(f"{SUMMARY_SHEET} repeats simulation group {key!r}")
        seen.add(key)
        formula_fields = (
            "dataset_batch_count", "episode_count", "frame_count", "duration",
            "scene_count", "first_date", "last_date",
        )
        _assert_summary_formula_cells(formula_ws, row_number, columns, formula_fields)
        actual_values = {
            "task_name": _clean_text(_summary_cached_value(cached_ws, row_number, columns["task_name"], "task_name")),
            "dataset_batch_count": _parse_nonnegative_int(
                _summary_cached_value(cached_ws, row_number, columns["dataset_batch_count"], "dataset_batch_count"),
                sheet=SUMMARY_SHEET, row=row_number, field="dataset_batch_count",
            ),
            "episode_count": _parse_nonnegative_int(
                _summary_cached_value(cached_ws, row_number, columns["episode_count"], "episode_count"),
                sheet=SUMMARY_SHEET, row=row_number, field="episode_count",
            ),
            "frame_count": _parse_nonnegative_int(
                _summary_cached_value(cached_ws, row_number, columns["frame_count"], "frame_count"),
                sheet=SUMMARY_SHEET, row=row_number, field="frame_count",
            ),
            "duration_seconds": duration_to_seconds(
                _summary_cached_value(cached_ws, row_number, columns["duration"], "duration"),
                sheet=SUMMARY_SHEET, row=row_number, field="duration",
            ),
            "scene_count": _parse_nonnegative_int(
                _summary_cached_value(cached_ws, row_number, columns["scene_count"], "scene_count"),
                sheet=SUMMARY_SHEET, row=row_number, field="scene_count",
            ),
            "generation_date_start": _parse_date(
                _summary_cached_value(cached_ws, row_number, columns["first_date"], "first_date"),
                sheet=SUMMARY_SHEET, row=row_number, field="first_date",
            ).isoformat(),
            "generation_date_end": _parse_date(
                _summary_cached_value(cached_ws, row_number, columns["last_date"], "last_date"),
                sheet=SUMMARY_SHEET, row=row_number, field="last_date",
            ).isoformat(),
        }
        for field, actual in actual_values.items():
            _compare_value(actual, expected[field], row=row_number, field=field)
    missing_groups = sorted(set(expected_rows) - seen)
    if missing_groups:
        raise DatasetStatisticsValidationError(
            f"{SUMMARY_SHEET} is missing simulation group row(s): {missing_groups!r}; "
            "add the corresponding formula rows before rebuilding JSON"
        )
    if total_row is None:
        raise DatasetStatisticsValidationError(f"{SUMMARY_SHEET} is missing 仿真总计")
    total_fields = ("dataset_batch_count", "episode_count", "frame_count", "duration", "scene_count")
    _assert_summary_formula_cells(formula_ws, total_row, columns, total_fields)
    totals = section["totals"]
    for field in ("dataset_batch_count", "episode_count", "frame_count", "scene_count"):
        actual = _parse_nonnegative_int(
            _summary_cached_value(cached_ws, total_row, columns[field], field),
            sheet=SUMMARY_SHEET, row=total_row, field=field,
        )
        _compare_value(actual, totals[field], row=total_row, field=field)
    duration = duration_to_seconds(
        _summary_cached_value(cached_ws, total_row, columns["duration"], "duration"),
        sheet=SUMMARY_SHEET, row=total_row, field="duration",
    )
    _compare_value(duration, totals["duration_seconds"], row=total_row, field="duration")
    _assert_summary_formula_cells(formula_ws, total_row, columns, ("first_date", "last_date"))
    if section["rows"]:
        first_date = _parse_date(
            _summary_cached_value(cached_ws, total_row, columns["first_date"], "first_date"),
            sheet=SUMMARY_SHEET,
            row=total_row,
            field="first_date",
        ).isoformat()
        last_date = _parse_date(
            _summary_cached_value(cached_ws, total_row, columns["last_date"], "last_date"),
            sheet=SUMMARY_SHEET,
            row=total_row,
            field="last_date",
        ).isoformat()
        _compare_value(
            first_date,
            min(str(item["generation_date_start"]) for item in section["rows"]),
            row=total_row,
            field="first_date",
        )
        _compare_value(
            last_date,
            max(str(item["generation_date_end"]) for item in section["rows"]),
            row=total_row,
            field="last_date",
        )


def _validate_summary(formula_ws: Any, cached_ws: Any, sections: Mapping[str, Any]) -> None:
    bounds = _summary_sections(formula_ws)
    for section_name in ("real_robot_teleoperation", "umi"):
        start, end = bounds[section_name]
        _reconcile_physical_summary(
            formula_ws, cached_ws, section_name, sections[section_name], start, end
        )
    start, end = bounds["simulation"]
    _reconcile_simulation_summary(
        formula_ws, cached_ws, sections["simulation"], start, end
    )


def build_statistics_payload(
    workbook_path: str | os.PathLike[str] = DEFAULT_WORKBOOK_PATH,
    *,
    validate_summary: bool = True,
) -> dict[str, Any]:
    """Read, validate, and aggregate the registry into the v1 JSON contract."""

    path = Path(workbook_path)
    if not path.is_file():
        raise DatasetStatisticsValidationError(f"Registry workbook does not exist: {path}")
    load_workbook, _ = _require_openpyxl()
    try:
        formula_book = load_workbook(path, data_only=False, read_only=False)
        cached_book = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:
        raise DatasetStatisticsValidationError(f"Unable to read registry workbook {path}: {exc}") from exc
    try:
        missing = [name for name in REQUIRED_SHEETS if name not in formula_book.sheetnames]
        if missing:
            raise DatasetStatisticsValidationError(
                f"Registry workbook is missing worksheet(s): {', '.join(missing)}"
            )
        _validate_no_excel_errors((formula_book, cached_book))
        dictionaries = _load_dictionaries(formula_book[DICTIONARY_SHEET])
        physical_rows = _validate_physical_rows(formula_book[PHYSICAL_SHEET], dictionaries)
        simulation_rows = _validate_simulation_rows(formula_book[SIMULATION_SHEET], dictionaries)
        _validate_unique_ids(physical_rows, simulation_rows)
        sections = _build_sections(physical_rows, simulation_rows)
        if validate_summary:
            _validate_summary(
                formula_book[SUMMARY_SHEET], cached_book[SUMMARY_SHEET], sections
            )
        return {
            "schema_version": SCHEMA_VERSION,
            "generated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
            **sections,
        }
    finally:
        formula_book.close()
        cached_book.close()


def _semantic_payload(payload: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "schema_version": payload.get("schema_version"),
        "real_robot_teleoperation": payload.get("real_robot_teleoperation"),
        "umi": payload.get("umi"),
        "simulation": payload.get("simulation"),
    }


def _load_existing_json(path: Path) -> dict[str, Any]:
    try:
        with path.open("r", encoding="utf-8") as handle:
            payload = json.load(handle)
    except FileNotFoundError as exc:
        raise DatasetStatisticsValidationError(f"Generated JSON does not exist: {path}") from exc
    except (OSError, json.JSONDecodeError) as exc:
        raise DatasetStatisticsValidationError(f"Generated JSON is invalid: {path}: {exc}") from exc
    if not isinstance(payload, dict) or payload.get("schema_version") != SCHEMA_VERSION:
        raise DatasetStatisticsValidationError(
            f"Generated JSON must use schema_version {SCHEMA_VERSION!r}"
        )
    required_keys = {
        "schema_version",
        "generated_at",
        "real_robot_teleoperation",
        "umi",
        "simulation",
    }
    if set(payload) != required_keys:
        raise DatasetStatisticsValidationError(
            f"Generated JSON must contain exactly: {', '.join(sorted(required_keys))}"
        )
    generated_at = payload.get("generated_at")
    if not isinstance(generated_at, str):
        raise DatasetStatisticsValidationError("Generated JSON generated_at must be an ISO timestamp")
    try:
        parsed_generated_at = datetime.fromisoformat(generated_at.replace("Z", "+00:00"))
    except ValueError as exc:
        raise DatasetStatisticsValidationError(
            "Generated JSON generated_at must be an ISO timestamp"
        ) from exc
    if parsed_generated_at.tzinfo is None:
        raise DatasetStatisticsValidationError(
            "Generated JSON generated_at must include a timezone"
        )
    return payload


def generate_statistics_json(
    workbook_path: str | os.PathLike[str] = DEFAULT_WORKBOOK_PATH,
    output_path: str | os.PathLike[str] = DEFAULT_JSON_PATH,
    *,
    check: bool = False,
    validate_summary: bool = True,
) -> BuildResult:
    """Validate the registry and atomically write (or check) generated JSON.

    Validation finishes before any temporary output is created.  Re-running the
    command against semantically identical workbook content preserves the JSON
    file byte-for-byte, including its original ``generated_at`` value.
    """

    payload = build_statistics_payload(workbook_path, validate_summary=validate_summary)
    output = Path(output_path)
    if check:
        existing = _load_existing_json(output)
        if _semantic_payload(existing) != _semantic_payload(payload):
            raise DatasetStatisticsValidationError(
                f"Generated JSON is out of sync with registry workbook: {output}"
            )
        payload["generated_at"] = existing.get("generated_at", payload["generated_at"])
        return BuildResult(payload=payload, changed=False)

    existing: dict[str, Any] | None = None
    if output.is_file():
        try:
            existing = _load_existing_json(output)
        except DatasetStatisticsValidationError:
            existing = None
    if existing is not None and _semantic_payload(existing) == _semantic_payload(payload):
        payload["generated_at"] = existing.get("generated_at", payload["generated_at"])
        return BuildResult(payload=payload, changed=False)

    output.parent.mkdir(parents=True, exist_ok=True)
    temp_name: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            newline="\n",
            dir=output.parent,
            prefix=f".{output.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temp_name = handle.name
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, output)
        temp_name = None
    finally:
        if temp_name is not None:
            try:
                os.unlink(temp_name)
            except FileNotFoundError:
                pass
    return BuildResult(payload=payload, changed=True)


__all__ = [
    "BuildResult",
    "DEFAULT_JSON_PATH",
    "DEFAULT_WORKBOOK_PATH",
    "DatasetStatisticsValidationError",
    "SCHEMA_VERSION",
    "build_statistics_payload",
    "duration_to_seconds",
    "generate_statistics_json",
]
